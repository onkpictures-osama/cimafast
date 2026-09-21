#!/usr/bin/env python3
"""Regression tests for the drama script analyzer.

Each case here corresponds to a defect that shipped and was reported as working:

  1. Props extraction used a bare substring search. The commit that claimed to
     "handle Arabic articles" added an OPTIONAL prefix with no word boundary,
     which matches identically with or without the alternation - a no-op. It
     read 'الدرجة' (step) as 'درج' (drawer) and 'الكوبري' (bridge) as 'كوب' (cup).
  2. Silent-character detection was handed the list of characters who HAVE
     dialogue, so it could only ever return speakers - the exact opposite - and
     its result was appended to a list those names were already in, making the
     whole block dead code. The 'silent_characters' key never existed.
  3. The analysis dashboard was rendered inside the confirm-import button, which
     then deleted parsed_script and called st.rerun(), discarding the frame it
     had just drawn. It was unreachable in the browser.

Run:  venv/bin/python test_script_parser.py
"""
import sys

from script_parser import (
    _extract_props_from_text,
    _detect_silent_characters,
    parse_script,
)

results = []


def check(name, condition, detail=""):
    results.append(bool(condition))
    print(f"{'✅' if condition else '❌'} {name}" + (f" → {detail}" if detail else ""))


def scene_from(script, **kw):
    out = parse_script("t.txt", script.encode("utf-8"), **kw)
    return out["scenes"][0] if out["scenes"] else {}


# ---------------------------------------------------------------- props
def test_props_word_boundaries():
    for text, forbidden, why in [
        ("وصل إلى الدرجة الأولى", "درج", "'الدرجة' (step) is not a drawer"),
        ("عبروا الكوبري بسرعة", "كوب", "'الكوبري' (bridge) is not a cup"),
        ("ركب المدرج", "درج", "'المدرج' (runway) is not a drawer"),
        ("سيفون الحمام", "سيف", "'سيفون' (cistern) is not a sword"),
    ]:
        got = _extract_props_from_text(text, [])
        check(f"no false positive: {why}", forbidden not in got, f"got {got}")


def test_props_still_found():
    for text, expected, why in [
        ("أحمد يمسك الكتاب", "كتاب", "ال prefix"),
        ("أخرج المسدس والمفتاح", "مفتاح", "وال prefix"),
        ("وضع الفنجان على الطاولة", "فنجان", "plain noun"),
        ("فتح الدرج", "درج", "genuine drawer still detected"),
    ]:
        got = _extract_props_from_text(text, [])
        check(f"still detects {expected} ({why})", expected in got, f"got {got}")


# ------------------------------------------------------- silent characters
def test_silent_excludes_speakers():
    # the original bug: a character WITH dialogue came back as "silent"
    silent = _detect_silent_characters("أحمد يدخل الغرفة", ["أحمد"], ["أحمد"])
    check("a speaking character is never silent", silent == [], f"got {silent}")


def test_silent_finds_non_speaker():
    silent = _detect_silent_characters("سعاد تقف في الركن", ["أحمد"], ["أحمد", "سعاد"])
    check("a non-speaking character in the action is silent", silent == ["سعاد"], f"got {silent}")


def test_silent_ignores_absent():
    silent = _detect_silent_characters("سعاد تقف في الركن", ["أحمد"], ["أحمد", "خالد"])
    check("a character absent from the scene is not silent", silent == [], f"got {silent}")


def test_silent_name_boundaries():
    # 'علي' must not match inside 'عليها'
    silent = _detect_silent_characters("وضع يده عليها", [], ["علي"])
    check("name not matched inside a longer word", silent == [], f"got {silent}")


# ----------------------------------------------------------- end to end
SCRIPT = """مشهد 1 - داخلي - نهار - شقة

أحمد يمسك الهاتف. سعاد تقف في الركن بلا كلام. وصل إلى الدرجة الأولى.

أحمد: أين المفتاح يا فاطمة؟

فاطمة: لا أعرف.
"""


def test_end_to_end():
    sc = scene_from(SCRIPT, known_characters=["سعاد", "خالد"])
    check("key 'silent_characters' exists (was never created before)",
          "silent_characters" in sc, f"keys={sorted(sc)[:6]}")
    check("سعاد detected as silent", sc.get("silent_characters") == ["سعاد"],
          f"got {sc.get('silent_characters')}")
    check("silent character added to the scene cast",
          "سعاد" in sc.get("characters", []), f"got {sc.get('characters')}")
    check("خالد (not in scene) excluded", "خالد" not in sc.get("characters", []))
    check("props clean end to end", sorted(sc.get("props", [])) == sorted(["هاتف", "مفتاح"]),
          f"got {sc.get('props')}")


def test_cross_scene_roster():
    """A character who speaks in one scene and is silent in another."""
    script = """مشهد 1 - داخلي - نهار - شقة

فاطمة تقف بجوار النافذة.

أحمد: أين أنت؟

مشهد 2 - داخلي - ليل - شقة

فاطمة: أنا هنا.
"""
    scenes = parse_script("t.txt", script.encode("utf-8"))["scenes"]
    check("two scenes parsed", len(scenes) == 2, f"got {len(scenes)}")
    if len(scenes) == 2:
        check("فاطمة silent in scene 1, speaking in scene 2",
              scenes[0].get("silent_characters") == ["فاطمة"]
              and "فاطمة" in scenes[1].get("characters", []),
              f"s1 silent={scenes[0].get('silent_characters')} s2 cast={scenes[1].get('characters')}")



# ------------------------------------------- docx tables reach the AI path
def test_docx_scene_tables_are_extracted():
    """Scene headers in Word tables must survive extraction.

    extract_lines originally kept only paragraphs, so every scene table -
    number, INT/EXT, day/night, location - was discarded before the text
    reached the AI. The model then had no numbers to copy and invented its own,
    which is exactly what was reported: "numbering is wrong, scenes skipped".
    """
    try:
        import docx as _docx
    except ImportError:
        check("python-docx available", False, "skipped")
        return
    from io import BytesIO
    from script_parser import extract_lines

    d = _docx.Document()
    for num, loc in (("مشهد 12", "شقة أحمد"), ("مشهد 13", "الشارع")):
        tbl = d.add_table(rows=1, cols=4)
        for cell, txt in zip(tbl.rows[0].cells, [num, loc, "نهار", "داخلي"]):
            cell.text = txt
        d.add_paragraph(f"وصف عند {loc}.")
    buf = BytesIO(); d.save(buf)

    lines = extract_lines("s.docx", buf.getvalue())
    joined = "\n".join(lines)
    check("scene 12 header survives extraction", "مشهد 12" in joined, joined[:60])
    check("scene 13 header survives extraction", "مشهد 13" in joined)
    check("original numbering preserved, not renumbered from 1",
          "مشهد 1 " not in joined and "مشهد 2 " not in joined)
    check("locations survive", "شقة أحمد" in joined and "الشارع" in joined)
    check("body paragraphs still present", "وصف عند" in joined)


# ------------------------------------------- refuse to analyse a non-script
def test_screenplay_detector():
    """A document with no screenplay markers must score near zero.

    The schema forces the model to return scenes, so it cannot answer "this is
    not a script" - it fabricates one instead. A real press release produced 13
    invented scenes and the system reported success. This is the gate that
    warns before any money is spent.
    """
    from script_parser import looks_like_screenplay

    press = ["شركة تطوير عقاري تطلق مشروعها الجديد في أكتوبر",
             "ويمثل المشروع الثاني للشركة في السوق المصرية على مساحة 80 فدانًا.",
             "وتبدأ أسعار الوحدات من 1.8 مليون جنيه بأنظمة سداد تصل إلى 10 سنوات."]
    conf, ev = looks_like_screenplay(press)
    check("press release scores ~0", conf < 0.3, f"score={conf} evidence={ev}")

    script = ["مشهد 1 - داخلي - نهار - شقة", "أحمد يجلس.", "أحمد: أهلا.",
              "فاطمة: أهلا بك.", "قطع",
              "مشهد 2 - خارجي - ليل - شارع", "سعاد تمشي."]
    conf2, ev2 = looks_like_screenplay(script)
    check("a real screenplay scores high", conf2 >= 0.8, f"score={conf2}")
    check("the two are clearly separated", conf2 - conf > 0.5, f"{conf} vs {conf2}")
    check("evidence is reported for the user to see",
          set(ev) == {"scene_headers", "int_ext", "day_night", "dialogue_lines", "cut_markers"})
    check("empty input does not crash", looks_like_screenplay([])[0] == 0.0)


if __name__ == "__main__":
    print("\n" + "=" * 60)
    print("Drama analyzer regression suite")
    print("=" * 60 + "\n")
    for fn in [test_props_word_boundaries, test_props_still_found,
               test_silent_excludes_speakers, test_silent_finds_non_speaker,
               test_silent_ignores_absent, test_silent_name_boundaries,
               test_end_to_end, test_cross_scene_roster,
               test_docx_scene_tables_are_extracted, test_screenplay_detector]:
        fn()
    passed, total = sum(results), len(results)
    print("\n" + "=" * 60)
    print(f"📊 {passed}/{total} tests passed")
    print("=" * 60 + "\n")
    sys.exit(0 if passed == total else 1)
