"""P10 الملابس — الغيارات، تعيين الغيار لكل مشهد، وقطع كل غيار.

    venv/bin/python tests/test_wardrobe.py

بيقفل: المظاهر القديمة بتترقّم غيارات من غير ما رقم موجود يتغيّر، المظهر
الجديد بياخد الغيار اللي بعده، تعيين الغيار بيرفض غيار شخصية تانية أو مشهد
مشروع تاني، "مشهد من غير غيار" بيتعد صح، القطع بتتحفظ بترتيبها والتكلفة
× النسخ، وكشف الملابس فيه كل غيار (حتى اللي مالوش قطع). والترقيم بيشتغل حتى
لو اليوزر "مشاهدة فقط".

قاعدة بيانات مؤقتة — عمرها ما بتلمس الإنتاج.
"""
from __future__ import annotations

import os
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
_TMP = tempfile.mkdtemp(prefix="cimafast-wardrobe-")
os.environ["STUDIO_DB_PATH"] = os.path.join(_TMP, "test.db")
os.environ.pop("DATABASE_URL", None)

import database  # noqa: E402

database.init_db()

import permissions  # noqa: E402
import repo  # noqa: E402
from database import fetch_all, run_query  # noqa: E402

_results = []


def test(fn):
    _results.append(fn)
    return fn


def _project(name):
    database.run_query("INSERT INTO companies (name) VALUES (?)", (name,))
    cid = fetch_all("SELECT MAX(id) AS id FROM companies")[0]["id"]
    run_query("INSERT INTO projects (name, project_type, company_id) VALUES (?, 'فيلم', ?)", (name, cid))
    return fetch_all("SELECT MAX(id) AS id FROM projects")[0]["id"]


P = _project("فيلم الملابس")
OTHER = _project("مشروع تاني")
NADIA = repo.add_character(P, "نادية", "بطل", "إنسان", "أنثى", "")
HASSAN = repo.add_character(P, "حسن", "مساعد", "إنسان", "ذكر", "")
STRANGER = repo.add_character(OTHER, "غريب", "بطل", "إنسان", "ذكر", "")
S1 = repo.add_scene(P, None, 1, "داخلي", "نهار", None, None, "")
S2 = repo.add_scene(P, None, 2, "خارجي", "ليل", None, None, "")
S3 = repo.add_scene(P, None, 3, "داخلي", "ليل", None, None, "")
S_OTHER = repo.add_scene(OTHER, None, 1, "داخلي", "نهار", None, None, "")
for sid, cid in ((S1, NADIA), (S1, HASSAN), (S2, NADIA), (S3, NADIA), (S3, HASSAN)):
    repo.link_character_to_scene(P, sid, cid)


def _changes(char_id):
    return [c for c in repo.wardrobe_changes(P) if c["character_id"] == char_id]


@test
def test_old_looks_get_numbered_without_moving_existing_numbers():
    # مظهر قديم من غير رقم (زي اللي اتعمل قبل P10)
    run_query("INSERT INTO character_looks (character_id, look_name, is_default) VALUES (?, 'قديم', 0)", (HASSAN,))
    run_query("UPDATE character_looks SET change_number=NULL WHERE character_id=? AND look_name='قديم'", (HASSAN,))
    assert repo.ensure_change_numbers(P) == 1
    nums = {c["look_name"]: c["change_number"] for c in _changes(HASSAN)}
    assert nums[repo.DEFAULT_LOOK_NAME] == 1 and nums["قديم"] == 2, nums
    assert repo.ensure_change_numbers(P) == 0


@test
def test_new_change_takes_next_number_and_labels():
    lid = repo.add_change(P, NADIA, "بدلة الفرح")
    ch = next(c for c in _changes(NADIA) if c["id"] == lid)
    assert ch["change_number"] == 2
    assert repo.change_label(ch) == "غيار 2 · بدلة الفرح"
    assert repo.change_label(_changes(NADIA)[0]) == "غيار 1"          # الاسم الافتراضي مابيتكتبش
    assert repo.add_change(P, STRANGER, "مش بتاعك") is None            # شخصية مشروع تاني


@test
def test_scene_assignment_is_scoped_and_counts_missing():
    n1, n2 = [c["id"] for c in _changes(NADIA)]
    other_char_look = _changes(HASSAN)[0]["id"]
    assert repo.scenes_missing_change(P) == 5
    done = repo.set_scene_changes(P, NADIA, {S1: n1, S2: n2, S3: other_char_look, S_OTHER: n1})
    assert done == 2, done                       # غيار حسن ومشهد المشروع التاني اتجاهلوا
    assert repo.scenes_missing_change(P) == 3
    m = repo.scene_change_map(P)
    assert m[(S1, NADIA)] == n1 and m[(S2, NADIA)] == n2 and (S3, NADIA) not in m
    repo.set_scene_changes(P, NADIA, {S2: None})
    assert (S2, NADIA) not in repo.scene_change_map(P)
    repo.set_scene_changes(P, NADIA, {S2: n2, S3: n2})
    rows = {r["id"]: r["look_id"] for r in repo.character_scenes_for_wardrobe(P, NADIA)}
    assert rows == {S1: n1, S2: n2, S3: n2}, rows


@test
def test_items_saved_in_order_with_cost_times_multiples():
    look = _changes(NADIA)[1]["id"]
    n = repo.save_change_items(P, look, [
        {"item_name": "بدلة سودا", "category": "بدلة", "cost": 3000, "multiples": 2, "status": "جاهز"},
        {"item_name": "  ", "cost": 99},                          # من غير اسم ← بيتشال
        {"item_name": "جزمة", "category": "جزمة", "cost": float("nan"), "multiples": None,
         "status": "محتاج شراء", "size": " 42 "},
    ])
    assert n == 2
    items = repo.items_of_change(P, look)
    assert [i["item_name"] for i in items] == ["بدلة سودا", "جزمة"]
    assert items[1]["multiples"] == 1 and items[1]["cost"] is None and items[1]["size"] == "42"
    ch = next(c for c in _changes(NADIA) if c["id"] == look)
    assert ch["items"] == 2 and ch["cost"] == 6000 and ch["not_ready"] == 1 and ch["scenes"] == 2
    assert repo.save_change_items(OTHER, look, [{"item_name": "x"}]) == 0     # مشروع غلط
    assert len(repo.items_of_change(P, look)) == 2


@test
def test_deleting_a_change_takes_its_items_and_assignments():
    lid = repo.add_change(P, HASSAN, "مؤقت")
    repo.save_change_items(P, lid, [{"item_name": "طاقية"}])
    repo.set_scene_changes(P, HASSAN, {S1: lid})
    repo.delete_character_look(P, lid)
    assert not fetch_all("SELECT 1 FROM wardrobe_items WHERE look_id=?", (lid,))
    assert not fetch_all("SELECT 1 FROM scene_character_looks WHERE look_id=?", (lid,))


@test
def test_numbering_works_for_view_only_users():
    run_query("INSERT INTO character_looks (character_id, look_name, is_default) VALUES (?, 'قديم 2', 0)", (NADIA,))
    run_query("UPDATE character_looks SET change_number=NULL WHERE look_name='قديم 2'")
    permissions.act_as("viewer")
    try:
        assert repo.ensure_change_numbers(P) == 1
    finally:
        permissions.act_as(None)


@test
def test_wardrobe_sheet_lists_every_change():
    import export
    from io import BytesIO
    from openpyxl import load_workbook
    data = export.build_wardrobe_sheet_excel(repo.project_by_id(P)[0], P, fetch_all)
    ws = load_workbook(BytesIO(data if isinstance(data, bytes) else data.getvalue())).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    flat = [" | ".join(str(v) for v in r if v is not None) for r in rows]
    suit = next(r for r in flat if "بدلة سودا" in r)
    assert "غيار 2 · بدلة الفرح" in suit and "2، 3" in suit, suit
    assert any("حسن" in r and "لسه مفيش قطع" in r for r in flat), flat[-6:]


def main():
    failed = 0
    for fn in _results:
        try:
            fn()
            print(f"  ok   {fn.__name__}")
        except Exception as exc:
            failed += 1
            print(f"  FAIL {fn.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n{len(_results) - failed}/{len(_results)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
