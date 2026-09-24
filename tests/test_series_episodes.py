"""المسلسل والحلقات — المرحلة الأولى (طلب المالك 2026-09-24).

    venv/bin/python tests/test_series_episodes.py

بيقفل: المسلسل مايتعملش من غير عدد حلقات والحلقات بتتعمل معاه، رقم المشهد
بيتكتب "حلقة/مشهد"، الاستيراد مابيعتبرش مشهد 1 في الحلقة 2 نسخة من مشهد 1
في الحلقة 1، والفيلم عمره ما بياخد رقم حلقة. ودفع الأرقام (رقم مستخدم) جوه
نفس الحلقة بس. وكمان إن قاعدة قديمة من غير scenes.episode_id بتتصلّح
بالمهاجرة (إضافة مشهد يدوي كانت بتقع على /v1 والإنتاج).

قاعدة بيانات مؤقتة — عمرها ما بتلمس الإنتاج.
"""
from __future__ import annotations

import os
import sqlite3
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
_TMP = tempfile.mkdtemp(prefix="cimafast-series-")
os.environ["STUDIO_DB_PATH"] = os.path.join(_TMP, "test.db")
os.environ.pop("DATABASE_URL", None)

import database  # noqa: E402

database.init_db()

import accounts  # noqa: E402
import auth  # noqa: E402
import repo  # noqa: E402
from database import fetch_all, run_query, scene_label  # noqa: E402
from importer import import_parsed_scenes  # noqa: E402

accounts.migrate_accounts({"melzayat": auth.hash_password("pw-123456", iterations=1000)})
CO, _ = accounts.create_company("melzayat", "شركة المسلسلات", "series_admin")
_results = []


def test(fn):
    _results.append(fn)
    return fn


def _raises(exc, fn, *a, **kw):
    try:
        fn(*a, **kw)
    except exc:
        return
    raise AssertionError(f"{fn.__name__} ماطلعش {exc.__name__}")


def _scene(n, episode=None, chars=("سلمى",)):
    sc = {"scene_number": n, "int_ext": "داخلي", "day_night": "نهار", "location_name": "شقة",
          "notes": "", "characters": list(chars), "props": []}
    if episode is not None:
        sc["episode_number"] = episode
    return sc


SERIES = accounts.create_project("series_admin", CO, "بلا أثر", "مسلسل", "1080p", "أفقي", "16:9",
                                 episode_count=10, type_details={"episode_minutes": 45})
FILM = accounts.create_project("series_admin", CO, "فيلم", "فيلم", "4K", "أفقي", "16:9")


@test
def test_series_needs_episode_count_and_gets_its_episodes():
    _raises(ValueError, accounts.create_project, "series_admin", CO, "من غير عدد", "مسلسل",
            "1080p", "أفقي", "16:9")
    _raises(ValueError, accounts.create_project, "series_admin", CO, "صفر", "مسلسل",
            "1080p", "أفقي", "16:9", episode_count=0)
    eps = [e["episode_number"] for e in repo.episode_overview(SERIES)]
    assert eps == list(range(1, 11)), eps
    assert repo.ensure_episodes(SERIES, 12) == 2 and repo.ensure_episodes(SERIES, 5) == 0
    assert not repo.episode_overview(FILM)
    details = fetch_all("SELECT type_details FROM projects WHERE id=?", (SERIES,))[0]["type_details"]
    assert '"episode_minutes": 45' in details


@test
def test_scene_label_carries_the_episode():
    assert scene_label({"scene_number": 12, "scene_suffix": None, "episode_number": 3}) == "3/12"
    assert scene_label({"scene_number": 12, "scene_suffix": "A", "episode_number": 3}) == "3/12A"
    assert scene_label({"scene_number": 12, "scene_suffix": None, "episode_number": None}) == "12"
    assert scene_label({"scene_number": 12}) == "12"


@test
def test_same_scene_numbers_in_two_episodes_are_two_scenes():
    s1 = import_parsed_scenes(SERIES, [_scene(1, 1), _scene(2, 1)], fetch_all, run_query)
    s2 = import_parsed_scenes(SERIES, [_scene(1, 2), _scene(2, 2), _scene(3, 2)], fetch_all, run_query)
    assert s1["scenes_added"] == 2 and s2["scenes_added"] == 3, (s1, s2)
    again = import_parsed_scenes(SERIES, [_scene(1, 2)], fetch_all, run_query)
    assert again["scenes_added"] == 0 and again["scenes_skipped"] == ["2/1"], again
    counts = {e["episode_number"]: e["scenes"] for e in repo.episode_overview(SERIES)}
    assert counts[1] == 2 and counts[2] == 3 and counts[3] == 0
    labels = [scene_label(s) for s in repo.scenes_of_project(SERIES)]
    assert labels == ["1/1", "1/2", "2/1", "2/2", "2/3"], labels


@test
def test_scenes_get_episode_id_and_unknown_episode_rows_are_created():
    import_parsed_scenes(SERIES, [_scene(1, 15)], fetch_all, run_query)
    row = fetch_all("SELECT s.episode_id, e.episode_number FROM scenes s JOIN episodes e ON e.id=s.episode_id "
                    "WHERE s.project_id=? AND s.episode_number=15", (SERIES,))
    assert row and row[0]["episode_number"] == 15


@test
def test_film_never_gets_episode_numbers():
    import_parsed_scenes(FILM, [_scene(1, 4), _scene(2)], fetch_all, run_query)
    eps = {s["episode_number"] for s in repo.scenes_of_project(FILM)}
    assert eps == {None}, eps
    assert not repo.episode_overview(FILM)


@test
def test_replace_episode_only_touches_that_episode():
    repo.delete_episode_scenes(SERIES, 2)
    counts = {e["episode_number"]: e["scenes"] for e in repo.episode_overview(SERIES)}
    assert counts[2] == 0 and counts[1] == 2
    repo.delete_episode_scenes(FILM, 1)            # مشروع تاني: مالهوش دعوة
    assert len(repo.scenes_of_project(FILM)) == 2
    import_parsed_scenes(SERIES, [_scene(1, 2), _scene(2, 2)], fetch_all, run_query)


@test
def test_manual_scene_takes_episode_and_shift_stays_inside_it():
    ep2 = next(e["id"] for e in repo.episode_overview(SERIES) if e["episode_number"] == 2)
    sid = repo.add_scene(SERIES, ep2, 1, "خارجي", "ليل", None, None, "")
    assert fetch_all("SELECT episode_number FROM scenes WHERE id=?", (sid,))[0]["episode_number"] == 2
    # رقم 1 مستخدم في الحلقة 2: الدفع يحرّك مشاهد الحلقة 2 بس
    repo.shift_scene_numbers_up_except(SERIES, 1, sid, 2)
    by_ep = {}
    for s in repo.scenes_of_project(SERIES):
        by_ep.setdefault(s["episode_number"], []).append(s["scene_number"])
    assert by_ep[1] == [1, 2], by_ep                # الحلقة 1 ماتحركتش
    assert sorted(by_ep[2]) == [1, 2, 3], by_ep
    assert not repo.other_scene_with_number(SERIES, 1, sid, 2)
    assert repo.other_scene_with_number(SERIES, 1, sid, 1)
    # فيلم: من غير حلقة، زي الأول بالظبط
    fid = repo.add_scene(FILM, None, 9, "داخلي", "نهار", None, None, "")
    assert fetch_all("SELECT episode_number FROM scenes WHERE id=?", (fid,))[0]["episode_number"] is None


@test
def test_move_scene_between_episodes():
    sid = fetch_all("SELECT id FROM scenes WHERE project_id=? AND episode_number=15", (SERIES,))[0]["id"]
    repo.set_scene_episode(SERIES, sid, 3)
    row = fetch_all("SELECT episode_number, episode_id FROM scenes WHERE id=?", (sid,))[0]
    ep3 = next(e["id"] for e in repo.episode_overview(SERIES) if e["episode_number"] == 3)
    assert row["episode_number"] == 3 and row["episode_id"] == ep3


@test
def test_character_sheet_lists_episode_scene_labels():
    import export
    from io import BytesIO
    from openpyxl import load_workbook
    data = export.build_characters_sheet_excel(repo.project_by_id(SERIES)[0], SERIES, fetch_all)
    ws = load_workbook(BytesIO(data if isinstance(data, bytes) else data.getvalue())).active
    row = next(r for r in ([c.value for c in row] for row in ws.iter_rows()) if "سلمى" in r)
    assert any(isinstance(v, str) and v.startswith("1/1، 1/2، 2/") for v in row), row


@test
def test_old_database_without_episode_id_gets_it():
    path = os.path.join(_TMP, "old.db")
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE projects (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, project_type TEXT)")
    con.execute("CREATE TABLE scenes (id INTEGER PRIMARY KEY AUTOINCREMENT, project_id INTEGER NOT NULL, "
                "scene_number INTEGER NOT NULL, int_ext TEXT, day_night TEXT, weather TEXT, "
                "location_variant_id INTEGER, notes TEXT)")
    con.commit()
    con.close()
    old = os.environ["STUDIO_DB_PATH"]
    os.environ["STUDIO_DB_PATH"] = path
    try:
        database.init_db()
        cols = {r["name"] for r in fetch_all("PRAGMA table_info(scenes)")}
        assert {"episode_id", "episode_number"} <= cols, cols
    finally:
        os.environ["STUDIO_DB_PATH"] = old


@test
def test_episode_number_from_filename():
    from views.import_tab import episode_from_filename as f
    assert f("الحلقة 07.docx") == 7 and f("حلقة_3.pdf") == 3 and f("Ep12 final.docx") == 12
    assert f("episode-٥.txt") == 5 and f("my_ep4.docx") == 4
    assert f("script.docx") is None and f("step3.docx") is None


def main():
    failed = 0
    for fn in _results:
        try:
            fn()
            print(f"  ok   {fn.__name__}")
        except Exception as exc:
            failed += 1
            print(f"  FAIL {fn.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n{len(_results) - failed}/{len(_results)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
