"""الممثل جوه باقي السيستم — رقم الكاست، التفريغ، الجدول، التتبع.

    venv/bin/python tests/test_casting_links.py

بيقفل: رقم الكاست فريد جوه المشروع وثابت (الترقيم التلقائي مبيحرّكش رقم
موجود)، اسم الممثل المتعاقد ورقمه بيوصلوا للـ strip بتاع الجدول وللـ DOOD
وللكشوفات، وتتبع أيام الشخصية بتواريخها.

قاعدة بيانات مؤقتة — عمرها ما بتلمس الإنتاج.
"""
from __future__ import annotations

import os
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
_TMP = tempfile.mkdtemp(prefix="cimafast-casting-")
os.environ["STUDIO_DB_PATH"] = os.path.join(_TMP, "test.db")
os.environ.pop("DATABASE_URL", None)

import database  # noqa: E402
import repo  # noqa: E402

database.init_db()
_results = []


def test(fn):
    _results.append(fn)
    return fn


def _company(name):
    database.run_query("INSERT INTO companies (name) VALUES (?)", (name,))
    return database.fetch_all("SELECT MAX(id) AS id FROM companies")[0]["id"]


def _project(company_id, name):
    database.run_query("INSERT INTO projects (name, project_type, company_id) VALUES (?, 'فيلم', ?)",
                       (name, company_id))
    return database.fetch_all("SELECT MAX(id) AS id FROM projects")[0]["id"]


def _scene(pid, number, *chars):
    sid = repo.add_scene(pid, None, number, "داخلي", "نهار", None, None, "")
    for c in chars:
        repo.link_character_to_scene(pid, sid, c)
    return sid


A = _company("شركة أ")
P = _project(A, "فيلم")
P2 = _project(A, "فيلم تاني")
HERO = repo.add_character(P, "سلمى", "بطل", "إنسان", "أنثى", "")
SIDE = repo.add_character(P, "سامي", "مساعد", "إنسان", "ذكر", "")
EXTRA = repo.add_character(P, "بواب", "كومبارس", "إنسان", "ذكر", "")
S1 = _scene(P, 1, HERO, SIDE)
S2 = _scene(P, 2, HERO)
S3 = _scene(P, 3, HERO, EXTRA)
S4 = _scene(P, 4, SIDE)
ACTOR = repo.add_actor({"full_name": "ممثلة تجريبية", "stage_name": "نجمة", "discoverable": 1},
                       owner_company_id=A, created_by="test")


@test
def test_auto_number_orders_by_scene_count_and_never_moves_existing():
    repo.set_cast_number(P, EXTRA, 7)            # رقم متحط بإيد قبل كده
    assert repo.auto_number_cast(P) == 2
    nums = {c["id"]: c["cast_number"] for c in repo.project_cast(P)}
    assert nums[EXTRA] == 7, "الترقيم التلقائي غيّر رقم كان متحط"
    assert nums[HERO] == 8 and nums[SIDE] == 9, nums   # سلمى 3 مشاهد قبل سامي 2
    assert repo.auto_number_cast(P) == 0


@test
def test_cast_number_is_unique_per_project_only():
    try:
        repo.set_cast_number(P, SIDE, 8)
    except repo.CastNumberTakenError:
        pass
    else:
        raise AssertionError("رقمين زي بعض في نفس المشروع")
    other = repo.add_character(P2, "حد تاني", "بطل", "إنسان", "ذكر", "")
    repo.set_cast_number(P2, other, 8)           # مشروع تاني: عادي
    repo.set_cast_number(P, EXTRA, None)
    repo.set_cast_number(P, EXTRA, 7)


@test
def test_cast_number_update_is_scoped_to_project():
    repo.set_cast_number(P2, HERO, 50)           # شخصية مش تبع P2 → صفر صفوف
    assert repo.cast_by_character(P)[HERO]["cast_number"] == 8


@test
def test_cast_actor_reaches_board_strip_and_dood():
    repo.cast_actor(ACTOR, P, HERO, "cast", "", "test")
    d1, d2 = repo.add_day(P), repo.add_day(P)
    days = repo.shooting_days(P)
    repo.update_day(P, days[0]["id"], shoot_date="2026-10-01")
    repo.update_day(P, days[1]["id"], shoot_date="2026-10-03")
    repo.save_layout(P, [{"day_id": days[0]["id"], "scene_ids": [S1, S2]},
                         {"day_id": days[1]["id"], "scene_ids": [S4]}])
    strip = next(s for s in repo.board_scenes(P) if s["id"] == S1)
    assert [c["num"] for c in strip["cast"]] == [8, 9], strip["cast"]
    assert strip["cast"][0]["actor"] == "نجمة"
    dood = repo.day_out_of_days(P)
    assert dood["dates"] == ["2026-10-01", "2026-10-03"]
    hero = next(r for r in dood["rows"] if r["character_id"] == HERO)
    assert hero["actor"] == "نجمة" and hero["cast_number"] == 8
    assert [r["cast_number"] for r in dood["rows"]] == [8, 9], "DOOD مش مترتب برقم الكاست"


@test
def test_tracking_counts_unscheduled_scenes_and_dates():
    t = repo.character_tracking(P, HERO)
    assert t["scene_count"] == 3 and t["scheduled_scenes"] == 2
    assert t["work_days"] == 1 and t["first_date"] == "2026-10-01"
    side = repo.character_tracking(P, SIDE)
    assert side["work_days"] == 2 and side["hold_days"] == 0
    assert [d["date"] for d in side["days"]] == ["2026-10-01", "2026-10-03"]
    assert repo.character_tracking(P, EXTRA)["work_days"] == 0


@test
def test_shortlist_shows_until_cast():
    other = repo.add_actor({"full_name": "مرشح", "discoverable": 1}, created_by="test")
    repo.cast_actor(other, P, SIDE, "shortlisted", "", "test")
    side = repo.cast_by_character(P)[SIDE]
    assert side["actor"] is None and [p["name"] for p in side["shortlist"]] == ["مرشح"]
    assert repo.cast_label(repo.cast_by_character(P)[HERO]) == "#8 سلمى — نجمة"


@test
def test_character_sheet_export_fills_actor_number_and_days():
    import export
    from openpyxl import load_workbook
    from io import BytesIO
    data = export.build_characters_sheet_excel(repo.project_by_id(P)[0], P, database.fetch_all)
    ws = load_workbook(BytesIO(data if isinstance(data, bytes) else data.getvalue())).active
    values = [[c.value for c in row] for row in ws.iter_rows()]
    hero = next(r for r in values if "سلمى" in r)
    assert 8 in hero and "نجمة" in hero and 1 in hero, hero   # رقم 8، الممثلة، يوم شغل واحد
    side = next(r for r in values if "سامي" in r)
    assert "مرشح (مرشح)" in side, side


def main():
    failed = 0
    for fn in _results:
        try:
            fn()
            print(f"  ok   {fn.__name__}")
        except Exception as exc:
            failed += 1
            print(f"  FAIL {fn.__name__}: {type(exc).__name__}: {exc}")
    print(f"\n{len(_results) - failed}/{len(_results)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
